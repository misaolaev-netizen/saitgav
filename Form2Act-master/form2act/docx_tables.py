from __future__ import annotations

import shutil
import tempfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from mailmerge import MailMerge

from form2act.template_scan import scan_template_tables


def _table_has_anchor(doc_path: Path, anchor: str) -> bool:
    with MailMerge(str(doc_path)) as doc:
        table, _idx, _row = doc._MailMerge__find_row_anchor(anchor)
        return table is not None


def append_tables_from_externals(
    main_template: Path,
    external_specs: list[dict],
) -> Path:
    if not external_specs:
        return main_template

    tmp = Path(tempfile.mkdtemp()) / "composed.docx"
    shutil.copy2(main_template, tmp)
    main_doc = Document(str(tmp))

    for spec in external_specs:
        ext = Path(spec["path"])
        if not ext.exists():
            continue
        anchor = spec.get("anchor")
        snip = Document(str(ext))
        src_tbl = None
        if anchor and _table_has_anchor(ext, anchor):
            with MailMerge(str(ext)) as m:
                table, _idx, _row = m._MailMerge__find_row_anchor(anchor)
                if table is not None:
                    src_tbl = table
        if src_tbl is None and snip.tables:
            src_tbl = snip.tables[0]._tbl
        if src_tbl is not None:
            main_doc.element.body.append(deepcopy(src_tbl))

    main_doc.save(str(tmp))
    return tmp


def list_external_tables_needed(tables: list[dict], main_path: str) -> list[dict]:
    main = Path(main_path)
    needed: list[dict] = []
    seen: set[str] = set()

    for t in tables:
        ext = t.get("source_template")
        if not ext or ext == main_path:
            continue
        key = f"{ext}:{t.get('anchor', '')}"
        if key in seen:
            continue
        anchor = t.get("anchor")
        if anchor and _table_has_anchor(main, anchor):
            continue
        if not anchor and ext:
            ext_anchors = scan_template_tables(Path(ext)).get("table_anchors") or []
            anchor = ext_anchors[0] if ext_anchors else None
        seen.add(key)
        needed.append({"path": ext, "anchor": anchor})

    return needed

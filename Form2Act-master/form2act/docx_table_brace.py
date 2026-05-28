from __future__ import annotations

import io
import re
from copy import deepcopy
from pathlib import Path

from docx import Document
from form2act.docx_edit import set_paragraph_text
from form2act.docx_placeholders import is_valid_merge_field_name
from form2act.merge_expand import PLACEHOLDER_RE, expand_merge_placeholders

_BRACE_IN_ROW_RE = re.compile(r"\{([^{}]+)\}")


def _row_brace_fields(row) -> list[str]:
    found: set[str] = set()
    for cell in row.cells:
        text = cell.text or ""
        for m in _BRACE_IN_ROW_RE.finditer(text):
            name = m.group(1).strip()
            if is_valid_merge_field_name(name):
                found.add(name)
    return sorted(found, key=str.casefold)


def _is_likely_data_row(fields: list[str]) -> bool:
    if len(fields) >= 2:
        return True
    if not fields:
        return False
    f = fields[0].casefold()
    return any(
        x in f
        for x in (
            "фамил",
            "fio",
            "групп",
            "n_pp",
            "имя",
            "отчеств",
        )
    )


def scan_brace_table_rows(template: Path) -> list[dict]:
    return scan_brace_table_rows_bytes(template.read_bytes())


def scan_brace_table_rows_bytes(docx_bytes: bytes) -> list[dict]:
    doc = Document(io.BytesIO(docx_bytes))
    results: list[dict] = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            fields = _row_brace_fields(row)
            if not _is_likely_data_row(fields):
                continue
            preferred = next(
                (f for f in fields if "фамил" in f.casefold() or "fio" in f.casefold()),
                fields[0],
            )
            results.append(
                {
                    "table_index": table_index,
                    "row_index": row_index,
                    "fields_in_row": fields,
                    "suggested_anchor": preferred,
                }
            )
    return results


def pick_primary_brace_table(tables: list[dict]) -> dict | None:
    if not tables:
        return None
    if len(tables) == 1:
        return tables[0]
    return max(tables, key=lambda t: (len(t["fields_in_row"]), -t["row_index"]))


def _substitute_braces_in_row(row, merge: dict[str, str]) -> None:
    expanded = expand_merge_placeholders(merge)

    def subst(text: str) -> str:
        def repl(match: re.Match[str]) -> str:
            key = match.group(1).strip()
            val = expanded.get(key)
            if val is None:
                for k, v in expanded.items():
                    if k.casefold() == key.casefold():
                        val = v
                        break
            if val is None or str(val).strip() == "":
                return match.group(0)
            return str(val)

        return PLACEHOLDER_RE.sub(repl, text)

    for cell in row.cells:
        for para in cell.paragraphs:
            raw = para.text or ""
            if "{" not in raw:
                continue
            new_text = subst(raw)
            if new_text != raw:
                set_paragraph_text(para, new_text)


def expand_brace_table_row(
    docx_bytes: bytes,
    *,
    table_index: int,
    row_index: int,
    rows_data: list[dict[str, str]],
) -> bytes:
    if not rows_data:
        return docx_bytes

    doc = Document(io.BytesIO(docx_bytes))
    if table_index >= len(doc.tables):
        raise ValueError(f"Таблица {table_index} не найдена в документе")
    table = doc.tables[table_index]
    if row_index >= len(table.rows):
        raise ValueError(f"Строка {row_index} не найдена в таблице")

    template_tr = table.rows[row_index]._tr
    for i in range(len(rows_data) - 1, 0, -1):
        new_tr = deepcopy(template_tr)
        table._tbl.insert(row_index + i, new_tr)

    for i, rd in enumerate(rows_data):
        _substitute_braces_in_row(table.rows[row_index + i], rd)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

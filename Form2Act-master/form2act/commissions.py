from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from form2act.config import DATA_DIR

COMMISSION_FILE = DATA_DIR / "commission.json"

_DATE_RE = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{2,4})")
_CHAIR_RE = re.compile(r"председатель", re.IGNORECASE)
_DEPUTY_RE = re.compile(r"заместитель\s+председател", re.IGNORECASE)
_SECRETARY_RE = re.compile(r"секретарь", re.IGNORECASE)


@dataclass
class Person:
    name: str = ""
    position: str = ""

    def as_line(self) -> str:
        name = self.name.strip()
        position = self.position.strip()
        if name and position:
            return f"{name} — {position}"
        return name or position


@dataclass
class Commission:
    title: str = "Состав государственной экзаменационной комиссии"
    specialty: str = ""
    chairman: Person = field(default_factory=Person)
    deputy_chairman: Person = field(default_factory=Person)
    members: list[Person] = field(default_factory=list)
    secretary: Person = field(default_factory=Person)


@dataclass
class DayMember:
    name: str
    role: str = "member"
    note: str = ""


@dataclass
class DaySchedule:
    sheet: str
    date: str = ""
    weekday: str = ""
    members: list[DayMember] = field(default_factory=list)


def _person_from(raw: Any) -> Person:
    if not isinstance(raw, dict):
        return Person()
    return Person(name=str(raw.get("name") or "").strip(), position=str(raw.get("position") or "").strip())


def load_commission(path: Path | None = None) -> Commission:
    path = path or COMMISSION_FILE
    if not path.exists():
        return Commission()
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return Commission()
    members = [_person_from(item) for item in raw.get("members") or [] if isinstance(item, dict)]
    return Commission(
        title=str(raw.get("title") or "Состав государственной экзаменационной комиссии").strip(),
        specialty=str(raw.get("specialty") or "").strip(),
        chairman=_person_from(raw.get("chairman")),
        deputy_chairman=_person_from(raw.get("deputy_chairman")),
        members=members,
        secretary=_person_from(raw.get("secretary")),
    )


def save_commission(commission: Commission, path: Path | None = None) -> None:
    path = path or COMMISSION_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    data = asdict(commission)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def commission_from_payload(payload: dict[str, Any]) -> Commission:
    base = load_commission()
    title = payload.get("title")
    specialty = payload.get("specialty")
    if isinstance(title, str) and title.strip():
        base.title = title.strip()
    if isinstance(specialty, str):
        base.specialty = specialty.strip()
    if "chairman" in payload:
        base.chairman = _person_from(payload.get("chairman"))
    if "deputy_chairman" in payload:
        base.deputy_chairman = _person_from(payload.get("deputy_chairman"))
    if "secretary" in payload:
        base.secretary = _person_from(payload.get("secretary"))
    if "members" in payload:
        members_payload = payload.get("members") or []
        if isinstance(members_payload, list):
            base.members = [_person_from(item) for item in members_payload if isinstance(item, dict)]
    return base


def _normalize_date(date: str) -> str:
    parts = date.split(".")
    if len(parts) != 3:
        return date
    day, month, year = parts
    if len(year) == 2:
        year = "20" + year
    return f"{day.zfill(2)}.{month.zfill(2)}.{year}"


def _split_date_weekday(value: str) -> tuple[str, str]:
    text = str(value or "").strip()
    if not text:
        return "", ""
    match = re.match(r"^(\d{1,2}\.\d{1,2}\.\d{2,4})\s*(?:\(([^)]*)\))?", text)
    if not match:
        return "", ""
    date = _normalize_date(match.group(1))
    weekday = (match.group(2) or "").strip()
    return date, weekday


def _classify_line(raw: str) -> tuple[str, str, str]:
    """Return (clean_name, role, note)."""
    line = raw.strip()
    role = "member"
    note = ""
    parts = re.split(r"\s+[-—–]\s+", line, maxsplit=1)
    if len(parts) == 2:
        name, tail = parts[0].strip(), parts[1].strip()
        if _DEPUTY_RE.search(tail):
            role = "deputy_chairman"
            note = tail
        elif _CHAIR_RE.search(tail):
            role = "chairman"
            note = tail
        elif _SECRETARY_RE.search(tail):
            role = "secretary"
            note = tail
        else:
            name = line
            note = ""
    else:
        if _DEPUTY_RE.search(line):
            role = "deputy_chairman"
        elif _CHAIR_RE.search(line):
            role = "chairman"
        elif _SECRETARY_RE.search(line):
            role = "secretary"
        name = line
    return name, role, note


def parse_schedule(path: Path) -> list[DaySchedule]:
    if not path.exists():
        return []
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return []

    result: list[DaySchedule] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        date, weekday = _split_date_weekday(sheet_name)
        if not date and rows[0]:
            first_cell = str(rows[0][0]) if rows[0][0] is not None else ""
            date, weekday = _split_date_weekday(first_cell)
        if not weekday and rows[0]:
            first_cell = str(rows[0][0]) if rows[0][0] is not None else ""
            _, alt_weekday = _split_date_weekday(first_cell)
            if alt_weekday:
                weekday = alt_weekday

        marker_row = -1
        for idx, row in enumerate(rows):
            text = " ".join(str(v) for v in row if v is not None).lower()
            if "комисси" in text:
                marker_row = idx
                break
        if marker_row < 0:
            continue

        members: list[DayMember] = []
        for row in rows[marker_row + 1 :]:
            cell = next((v for v in row if v is not None and str(v).strip()), None)
            if cell is None:
                continue
            text = str(cell).strip()
            if not text or text.lower().startswith("комисси"):
                continue
            name, role, note = _classify_line(text)
            members.append(DayMember(name=name, role=role, note=note))

        if not members:
            continue

        result.append(DaySchedule(sheet=sheet_name, date=date, weekday=weekday, members=members))

    return result


def _find_schedule_file() -> Path | None:
    for path in sorted(DATA_DIR.glob("*.xlsx")):
        name = path.name
        if name.startswith("~$"):
            continue
        if name.casefold().startswith("защита"):
            return path
    return None


def commission_payload() -> dict[str, Any]:
    commission = load_commission()
    schedule_path = _find_schedule_file()
    schedule = parse_schedule(schedule_path) if schedule_path else []
    return {
        "commission": asdict(commission),
        "schedule": [asdict(day) for day in schedule],
        "schedule_file": schedule_path.name if schedule_path else None,
    }

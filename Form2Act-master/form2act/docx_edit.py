from __future__ import annotations

import html as html_lib
import io
import re
from html.parser import HTMLParser

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

_MERGE_FIELD_RE = re.compile(r"\{([^{}]+)\}")


def _paragraphs_in_container(parent) -> list[Paragraph]:
    result: list[Paragraph] = []
    if hasattr(parent, "element") and hasattr(parent.element, "body"):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        return result

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            result.append(Paragraph(child, parent))
        elif isinstance(child, CT_Tbl):
            table = Table(child, parent)
            for row in table.rows:
                for cell in row.cells:
                    result.extend(_paragraphs_in_container(cell))
    return result


def paragraph_texts_from_docx_bytes(docx_bytes: bytes) -> list[str]:
    doc = Document(io.BytesIO(docx_bytes))
    return [p.text for p in _paragraphs_in_container(doc)]


def set_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


class _HtmlTextExtractor(HTMLParser):
    _BLOCK = frozenset(("p", "td", "th", "li"))

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.texts: list[str] = []
        self._depth = 0
        self._buf: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        tag = tag.lower()
        if tag in self._BLOCK:
            if self._depth == 0:
                self._buf = []
            self._depth += 1

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag not in self._BLOCK or self._depth == 0:
            return
        self._depth -= 1
        if self._depth == 0:
            text = html_lib.unescape("".join(self._buf))
            text = re.sub(r"\s+", " ", text).strip()
            self.texts.append(text)
            self._buf = []

    def handle_data(self, data: str) -> None:
        if self._depth > 0:
            self._buf.append(data)


def extract_texts_from_preview_html(html: str) -> list[str]:
    parser = _HtmlTextExtractor()
    parser.feed(html)
    parser.close()
    return parser.texts


def extract_texts_from_docx_preview_html(html: str) -> list[str]:
    texts: list[str] = []
    for section in re.finditer(
        r'<section[^>]*\bclass="[^"]*\bdocx\b[^"]*"[^>]*>([\s\S]*?)</section>',
        html,
        re.I,
    ):
        chunk = section.group(1)
        for m in re.finditer(r"<(p|td|th|li)\b[^>]*>([\s\S]*?)</\1>", chunk, re.I):
            inner = re.sub(r"<[^>]+>", " ", m.group(2))
            t = html_lib.unescape(re.sub(r"\s+", " ", inner).strip())
            texts.append(t)
    if texts:
        return texts
    return extract_texts_from_preview_html(html)


def _baseline_matches_docx(docx_bytes: bytes, baseline: list[str]) -> bool:
    if not baseline:
        return False
    actual = paragraph_texts_from_docx_bytes(docx_bytes)
    if len(baseline) > len(actual):
        return False
    for expected, current in zip(baseline, actual):
        if (expected or "").strip() != (current or "").strip():
            return False
    return True


def _layouts_compatible(docx_baseline: list[str], edited: list[str]) -> bool:
    if len(docx_baseline) != len(edited) or not docx_baseline:
        return False
    matches = 0
    for a, b in zip(docx_baseline, edited):
        at = (a or "").strip()
        bt = (b or "").strip()
        if not at or not bt:
            matches += 1
            continue
        if at == bt or at in bt or bt in at:
            matches += 1
        elif len(at) >= 20 and len(bt) >= 20 and at[:20] == bt[:20]:
            matches += 1
    return matches >= max(3, int(len(docx_baseline) * 0.2))


def apply_saved_html_edits(
    docx_bytes: bytes,
    html: str | None,
    *,
    stored_baseline: list[str] | None = None,
) -> bytes:
    if not html or not str(html).strip():
        return docx_bytes
    docx_baseline = paragraph_texts_from_docx_bytes(docx_bytes)
    edited = extract_texts_from_docx_preview_html(html)
    if not edited or not docx_baseline:
        return docx_bytes

    if stored_baseline:
        if not _baseline_matches_docx(docx_bytes, stored_baseline):
            return docx_bytes
        baseline = stored_baseline
    else:
        if len(edited) != len(docx_baseline):
            return docx_bytes
        if not _layouts_compatible(docx_baseline, edited):
            return docx_bytes
        baseline = docx_baseline

    if len(edited) != len(baseline):
        return docx_bytes
    return apply_edits_to_docx(docx_bytes, baseline, edited)


def extract_merge_overrides_from_html(html: str) -> dict[str, str]:
    overrides: dict[str, str] = {}
    for match in _MERGE_FIELD_RE.finditer(html):
        key = match.group(1).strip()
        if key:
            overrides.setdefault(key, "")
    for match in re.finditer(
        r'data-merge-field="([^"]+)"[^>]*>(.*?)</span>',
        html,
        re.I | re.S,
    ):
        key = html_lib.unescape(match.group(1).strip())
        val = re.sub(r"<[^>]+>", "", match.group(2))
        val = html_lib.unescape(re.sub(r"\s+", " ", val).strip())
        if key:
            overrides[key] = val
    return overrides


def apply_edits_to_docx(
    docx_bytes: bytes,
    baseline: list[str] | None,
    edited: list[str] | None,
) -> bytes:
    if not baseline or not edited:
        return docx_bytes
    if not _baseline_matches_docx(docx_bytes, baseline):
        return docx_bytes

    doc = Document(io.BytesIO(docx_bytes))
    paragraphs = _paragraphs_in_container(doc)
    if not paragraphs:
        return docx_bytes

    if len(baseline) == len(edited) == len(paragraphs):
        for para, old, new in zip(paragraphs, baseline, edited):
            new_t = (new or "").strip()
            old_t = (old or "").strip()
            if new_t and new_t != old_t and new_t != para.text.strip():
                set_paragraph_text(para, new_t)
    elif len(baseline) == len(edited) and len(baseline) <= len(paragraphs):
        for para, old, new in zip(paragraphs[: len(baseline)], baseline, edited):
            new_t = (new or "").strip()
            old_t = (old or "").strip()
            if new_t and new_t != old_t and new_t != para.text.strip():
                set_paragraph_text(para, new_t)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

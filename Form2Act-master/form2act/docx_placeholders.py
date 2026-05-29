from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

from form2act.docx_edit import _paragraphs_in_container, set_paragraph_text
from form2act.merge_expand import PLACEHOLDER_RE, expand_merge_placeholders

_BRACE_SCAN_RE = re.compile(r"\{([^{}]+)\}")
_MAILMERGE_XML_RE = re.compile(
    r"MERGEFIELD\s+([^\s\\]+)|MERGEFIELD\s+\\?\"([^\"\\]+)\"",
    re.IGNORECASE,
)
_GUID_FIELD_RE = re.compile(
    r"^[0-9A-Fa-f]{8}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{4}-[0-9A-Fa-f]{12}$"
)


def is_valid_merge_field_name(name: str) -> bool:
    n = (name or "").strip()
    if not n or len(n) > 80:
        return False
    if _GUID_FIELD_RE.match(n):
        return False
    if n.startswith("\\") or "xmlns" in n.casefold():
        return False
    if re.fullmatch(r"[0-9A-Fa-f\-]{20,}", n):
        return False
    return True


def _filter_field_names(names: set[str] | list[str]) -> list[str]:
    return sorted({n.strip() for n in names if is_valid_merge_field_name(n)}, key=str.casefold)


_MERGE_FIELD_ALIASES: dict[str, str] = {
    "дата защиты": "Дата_защиты_диплома",
    "дата защиты диплома": "Дата_защиты_диплома",
    "дата_защиты": "Дата_защиты_диплома",
    "датазащиты": "Дата_защиты_диплома",
    "дата дэ": "Дата_ДЭ",
    "фио": "Фамилия_имя_отчество",
}


def _normalize_field_key(key: str) -> str:
    k = (key or "").strip()
    if not k:
        return k
    alias = _MERGE_FIELD_ALIASES.get(k.casefold().replace("_", " "))
    if alias:
        return alias
    alias = _MERGE_FIELD_ALIASES.get(re.sub(r"\s+", " ", k.casefold()))
    return alias or k


def _is_merge_value_empty(value: str, field_name: str) -> bool:
    s = (value or "").strip()
    if not s:
        return True
    if s == f"{{{field_name}}}":
        return True
    if re.fullmatch(r"[\-—–]+", s):
        return True
    return False


def _merge_lookup(merge: dict[str, str], key: str) -> str:
    key = _normalize_field_key(key)
    if key in merge:
        return str(merge[key] or "")
    key_l = key.casefold()
    for k, v in merge.items():
        if k.casefold() == key_l:
            return str(v or "")
    return ""


def _scan_xml_parts(docx: zipfile.ZipFile) -> tuple[set[str], set[str]]:
    mailmerge: set[str] = set()
    brace: set[str] = set()
    for name in docx.namelist():
        if not name.startswith("word/") or not name.endswith(".xml"):
            continue
        try:
            xml = docx.read(name).decode("utf-8", errors="ignore")
        except KeyError:
            continue
        for m in _MAILMERGE_XML_RE.finditer(xml):
            field = (m.group(1) or m.group(2) or "").strip()
            if is_valid_merge_field_name(field):
                mailmerge.add(field)
        for m in _BRACE_SCAN_RE.finditer(xml):
            field = _normalize_field_key(m.group(1).strip())
            if is_valid_merge_field_name(field):
                brace.add(field)
        plain = re.sub(r"<[^>]+>", "", xml)
        for m in _BRACE_SCAN_RE.finditer(plain):
            field = _normalize_field_key(m.group(1).strip())
            if is_valid_merge_field_name(field):
                brace.add(field)
        for block in re.findall(r"<w:p\b[^>]*>.*?</w:p>", xml, flags=re.DOTALL):
            para_text = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", block))
            for m in _BRACE_SCAN_RE.finditer(para_text):
                field = _normalize_field_key(m.group(1).strip())
                if is_valid_merge_field_name(field):
                    brace.add(field)
    return mailmerge, brace


def scan_fields_from_docx_bytes(docx_bytes: bytes) -> tuple[list[str], list[str], list[str]]:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        mailmerge, brace = _scan_xml_parts(zf)
    doc = Document(io.BytesIO(docx_bytes))
    for para in _paragraphs_everywhere(doc):
        for m in _BRACE_SCAN_RE.finditer(para.text or ""):
            name = _normalize_field_key(m.group(1).strip())
            if is_valid_merge_field_name(name):
                brace.add(name)
    mailmerge_list = _filter_field_names(mailmerge)
    brace_list = _filter_field_names(brace)
    all_names = sorted(set(mailmerge_list) | set(brace_list), key=str.casefold)
    return mailmerge_list, brace_list, all_names


def scan_fields_from_template(template: Path) -> tuple[list[str], list[str], list[str]]:
    return scan_fields_from_docx_bytes(template.read_bytes())


def _paragraphs_everywhere(doc: Document) -> list[Paragraph]:
    paras = list(_paragraphs_in_container(doc))
    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if hf is None:
                continue
            paras.extend(hf.paragraphs)
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paras.extend(cell.paragraphs)
    return paras


def brace_fields_in_template(template: Path) -> list[str]:
    _mailmerge, brace, _all = scan_fields_from_template(template)
    return brace


def brace_fields_in_docx_bytes(docx_bytes: bytes) -> list[str]:
    _mailmerge, brace, _all = scan_fields_from_docx_bytes(docx_bytes)
    return brace


def template_merge_field_names(template: Path) -> tuple[list[str], list[str], list[str]]:
    if not template.exists():
        return [], [], []
    mailmerge, brace, all_names = scan_fields_from_template(template)
    try:
        from form2act.custom_fields import word_fields_in_template

        mm_lib = word_fields_in_template(template)
        mailmerge = _filter_field_names(set(mailmerge) | set(mm_lib))
        all_names = sorted(set(mailmerge) | set(brace), key=str.casefold)
    except Exception:
        pass
    return mailmerge, brace, all_names


def replace_brace_placeholders_in_docx_bytes(
    docx_bytes: bytes,
    merge_data: dict[str, str],
) -> bytes:
    expanded = expand_merge_placeholders(merge_data)

    def subst(text: str) -> str:
        def repl(match: re.Match[str]) -> str:
            key = match.group(1).strip()
            val = _merge_lookup(expanded, key).strip()
            if not val:
                return match.group(0)
            return val

        return PLACEHOLDER_RE.sub(repl, text)

    doc = Document(io.BytesIO(docx_bytes))
    for para in _paragraphs_everywhere(doc):
        raw = para.text or ""
        if "{" not in raw:
            continue
        new_text = subst(raw)
        if new_text != raw:
            set_paragraph_text(para, new_text)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def missing_template_fields(
    template: Path,
    merge_data: dict[str, str],
) -> tuple[list[str], list[str], list[str]]:
    mailmerge, brace, all_names = template_merge_field_names(template)
    expanded = expand_merge_placeholders(merge_data)
    missing: list[str] = []
    missing_brace: list[str] = []
    for name in all_names:
        val = _merge_lookup(expanded, name)
        if _is_merge_value_empty(val, name):
            missing.append(name)
            if name in brace:
                missing_brace.append(name)
    return all_names, missing, missing_brace


def fill_status_for_all_students(
    template: Path,
    records: dict,
    get_merge_payload,
    *,
    fields_subset: set[str] | None = None,
) -> dict:
    from form2act.merge_expand import expand_merge_placeholders

    all_fields, _, _ = template_merge_field_names(template)
    if fields_subset is not None:
        all_fields = [f for f in all_fields if f in fields_subset]
    by_id: dict[str, dict] = {}
    field_gap_count: dict[str, int] = {f: 0 for f in all_fields}

    for sid, rec in list(records.items()):
        merge = expand_merge_placeholders(get_merge_payload(sid))
        _all, missing, _ = missing_template_fields(template, merge)
        if fields_subset is not None:
            missing = [m for m in missing if m in fields_subset]
        if missing:
            by_id[sid] = {
                "id": sid,
                "fio": rec.get("_fio", sid),
                "missing": missing,
                "missing_count": len(missing),
            }
            for f in missing:
                field_gap_count[f] = field_gap_count.get(f, 0) + 1

    gaps_sorted = sorted(
        [f for f, n in field_gap_count.items() if n > 0],
        key=lambda f: (-field_gap_count[f], f.casefold()),
    )
    return {
        "template_fields": all_fields,
        "students_with_gaps": len(by_id),
        "total_students": len(records),
        "by_id": by_id,
        "fields_often_missing": gaps_sorted,
        "field_gap_count": field_gap_count,
    }
